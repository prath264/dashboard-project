from io import BytesIO
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ExportSection = tuple[str, list[tuple[str, str]], list[dict[str, Any]]]


def _excel_value(value: Any) -> Any:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def make_excel(sections: list[ExportSection]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, columns, rows in sections:
        sheet = workbook.create_sheet(title[:31])
        sheet.append([label for label, _ in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for row in rows:
            sheet.append([_excel_value(row.get(key)) for _, key in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 12), 40)
            sheet.column_dimensions[cells[0].column_letter].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pdf(sections: list[ExportSection], title: str) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.35 * inch,
        leftMargin=0.35 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch,
    )
    styles = getSampleStyleSheet()
    styles["Title"].alignment = TA_CENTER
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for section_title, columns, rows in sections:
        story.append(Paragraph(section_title, styles["Heading2"]))
        table_data = [[label for label, _ in columns]]
        table_data.extend([
            [str(row.get(key) if row.get(key) is not None else "") for _, key in columns]
            for row in rows
        ])
        page_width = landscape(letter)[0] - document.leftMargin - document.rightMargin
        table = Table(
            table_data,
            colWidths=[page_width / len(columns)] * len(columns),
            repeatRows=1,
            hAlign="LEFT",
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("LEADING", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend([table, Spacer(1, 12)])
    document.build(story)
    return output.getvalue()
